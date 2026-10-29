"""Tests for get_cord_stats.py against a small workbook. Needs openpyxl.

Run with:
  python3 -m unittest discover -s bin -p 'test_*.py'
"""

import contextlib
import datetime
import io
import json
import os
import shutil
import sys
import tempfile
import unittest

try:
    import openpyxl
except ImportError:
    openpyxl = None

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
if openpyxl:
    import get_cord_stats  # noqa: E402

HEADER = ["LAST_UPDATED", "UPDATE_REASON", "SOURCE", "GEO", "RECORD_COUNT",
          "PERSON_COUNT", "NAME_FEATURES", "DOB_FEATURES"]
OLD_DATE = datetime.datetime(2025, 1, 2)
DATE = datetime.datetime(2026, 10, 29)
ROWS = [
    [OLD_DATE, "manual note", "ofac", "london", 7, 7, 7, 3],
    [OLD_DATE, "keep me", "gleif", "london", 5, 0, 5, 0],
    [OLD_DATE, "other", "ofac", "moscow", 9, 9, 9, 9],
]
RECORDS = [
    {"DATA_SOURCE": "OFAC", "RECORD_ID": "1", "RECORD_TYPE": "PERSON", "NAME_FULL": "Bob Smith"},
    {"DATA_SOURCE": "OFAC", "RECORD_ID": "2", "RECORD_TYPE": "PERSON", "NAME_FULL": "Al Jones"},
]


@unittest.skipUnless(openpyxl, "openpyxl is not installed: pip install openpyxl")
class GetCordStatsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp)
        self.stats = os.path.join(self.tmp, "CORD_stats.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "file_stats"
        for row in [HEADER] + ROWS:
            ws.append(row)
        ws.cell(row=4, column=5).font = openpyxl.styles.Font(bold=True)
        wb.save(self.stats)

    def write_jsonl(self, name, records):
        path = os.path.join(self.tmp, name)
        with open(path, "w") as file:
            file.writelines(json.dumps(record) + "\n" for record in records)
        return path

    def run_stats(self, *args):
        out = io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(out):
            code = get_cord_stats.main(["--stats", self.stats, "--date", "2026-10-29", *args])
        return code, out.getvalue()

    def rows(self):
        return [list(row) for row in openpyxl.load_workbook(self.stats).worksheets[0].values]

    def test_updates_only_the_given_files_row(self):
        path = self.write_jsonl("ofac-london.jsonl", RECORDS)
        code, output = self.run_stats("--reason", "Changed in 0.2.0", path)
        self.assertEqual(code, 0, output)
        rows = self.rows()
        self.assertEqual(rows[0], HEADER)
        # Counts rewritten, DOB no longer present so 0, date and reason set.
        self.assertEqual(rows[1], [DATE, "Changed in 0.2.0", "ofac", "london", 2, 2, 2, 0])
        self.assertEqual(rows[2:], ROWS[1:])

    def test_new_file_gets_a_new_formatted_row(self):
        path = self.write_jsonl("icij-lasvegas.jsonl", RECORDS[:1])
        code, output = self.run_stats("--reason", "Added in 0.2.0", path)
        self.assertEqual(code, 0, output)
        rows = self.rows()
        self.assertEqual(rows[1:4], ROWS)
        self.assertEqual(rows[4], [DATE, "Added in 0.2.0", "icij", "lasvegas", 1, 1, 1, 0])
        ws = openpyxl.load_workbook(self.stats).worksheets[0]
        self.assertTrue(ws.cell(row=5, column=5).font.bold)

    def test_name_splits_on_the_last_hyphen(self):
        self.assertEqual(get_cord_stats.source_and_geo("/x/open-sanctions-london.jsonl"),
                         ("open-sanctions", "london"))

    def test_name_without_hyphen_stops_before_writing(self):
        before = os.path.getmtime(self.stats)
        path = self.write_jsonl("bob.jsonl", RECORDS)
        code, output = self.run_stats("--reason", "Added in 0.2.0", path)
        self.assertEqual(code, 2)
        self.assertIn("bob.jsonl: expected <source>-<geo>.jsonl", output)
        self.assertEqual(os.path.getmtime(self.stats), before)


if __name__ == "__main__":
    unittest.main()
